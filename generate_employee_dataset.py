"""Generate a synthetic employee dataset and write it to an Excel workbook."""

from __future__ import annotations

import random
import unicodedata
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from faker import Faker
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROW_COUNT = 150
SEED = 42
OUTPUT_PATH = Path(__file__).parent / "employee_dataset.xlsx"
SHEET_NAME = "Employees"
EMAIL_DOMAIN = "example.com"
ACTIVE_RATE = 0.85
HIRE_START = date(2015, 1, 1)
HIRE_END = date.today()
MAX_COLUMN_WIDTH = 40

COLUMNS = [
    "employee_id",
    "first_name",
    "last_name",
    "email",
    "department",
    "job_title",
    "hire_date",
    "salary",
    "city",
    "is_active",
]

DEPARTMENT_TITLES = {
    "Engineering": [
        "Software Engineer",
        "Senior Software Engineer",
        "QA Engineer",
        "DevOps Engineer",
        "Engineering Manager",
    ],
    "Sales": [
        "Sales Representative",
        "Account Executive",
        "Solutions Consultant",
        "Sales Manager",
    ],
    "Marketing": [
        "Marketing Specialist",
        "Content Strategist",
        "SEO Analyst",
        "Marketing Manager",
    ],
    "HR": [
        "Recruiter",
        "HR Generalist",
        "People Operations Specialist",
        "HR Manager",
    ],
    "Finance": [
        "Financial Analyst",
        "Accountant",
        "Payroll Specialist",
        "Finance Manager",
    ],
    "Operations": [
        "Operations Analyst",
        "Logistics Coordinator",
        "Supply Chain Planner",
        "Operations Manager",
    ],
    "Support": [
        "Support Specialist",
        "Technical Support Engineer",
        "Customer Success Manager",
        "Support Team Lead",
    ],
}

SALARY_RANGES = {
    "Engineering": (85_000, 180_000),
    "Sales": (55_000, 140_000),
    "Marketing": (55_000, 120_000),
    "HR": (50_000, 110_000),
    "Finance": (60_000, 135_000),
    "Operations": (52_000, 115_000),
    "Support": (45_000, 95_000),
}

NUMBER_FORMATS = {"hire_date": "yyyy-mm-dd", "salary": "#,##0"}


def _slugify(value: str) -> str:
    """Strip accents and punctuation so a name is safe inside an email local part."""
    folded = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return "".join(char for char in folded if char.isalnum()).lower()


def _build_email(first_name: str, last_name: str, taken: set[str]) -> str:
    base = f"{_slugify(first_name)}.{_slugify(last_name)}"
    local_part = base
    suffix = 2
    while local_part in taken:
        local_part = f"{base}{suffix}"
        suffix += 1
    taken.add(local_part)
    return f"{local_part}@{EMAIL_DOMAIN}"


def _random_hire_date() -> date:
    return HIRE_START + timedelta(days=random.randint(0, (HIRE_END - HIRE_START).days))


def build_dataframe(row_count: int = ROW_COUNT) -> pd.DataFrame:
    random.seed(SEED)
    Faker.seed(SEED)
    fake = Faker("en_US")

    departments = list(DEPARTMENT_TITLES)
    taken_emails: set[str] = set()
    records = []

    for index in range(1, row_count + 1):
        first_name = fake.first_name()
        last_name = fake.last_name()
        department = random.choice(departments)
        lowest, highest = SALARY_RANGES[department]
        records.append(
            {
                "employee_id": f"EMP{index:04d}",
                "first_name": first_name,
                "last_name": last_name,
                "email": _build_email(first_name, last_name, taken_emails),
                "department": department,
                "job_title": random.choice(DEPARTMENT_TITLES[department]),
                "hire_date": _random_hire_date(),
                "salary": round(random.randint(lowest, highest) / 500) * 500,
                "city": fake.city(),
                "is_active": random.random() < ACTIVE_RATE,
            }
        )

    return pd.DataFrame(records, columns=COLUMNS)


def _format_sheet(sheet: Worksheet, frame: pd.DataFrame) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for position, name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(position)
        widest = max([len(name), *(len(str(value)) for value in frame[name])])
        sheet.column_dimensions[letter].width = min(widest + 2, MAX_COLUMN_WIDTH)

        number_format = NUMBER_FORMATS.get(name)
        if number_format is None:
            continue
        for (cell,) in sheet.iter_rows(min_row=2, min_col=position, max_col=position):
            cell.number_format = number_format


def write_workbook(frame: pd.DataFrame, output_path: Path = OUTPUT_PATH) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=SHEET_NAME, index=False)
        _format_sheet(writer.sheets[SHEET_NAME], frame)


def main() -> None:
    frame = build_dataframe()
    write_workbook(frame)
    print(f"Wrote {len(frame)} rows x {len(frame.columns)} columns to {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    main()
